#!/usr/bin/env python3
"""Generate the MEDAGENT-X v2 pipeline Word documentation."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

OUTPUT_PATH = (
    Path(__file__).resolve().parents[1]
    / "docs"
    / "MEDAGENT-X-v2-Complete-Pipeline-Documentation.docx"
)


def add_title(doc: Document, text: str) -> None:
    doc.add_heading(text, level=0)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_para(doc: Document, text: str) -> None:
    doc.add_paragraph(text)


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def add_numbered(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def build_document() -> Document:
    doc = Document()

    # Title page
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("MEDAGENT-X v2\nComplete Pipeline Documentation")
    run.bold = True
    run.font.size = Pt(24)
    doc.add_paragraph()
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(
        "From DICOM acquisition through vision, retrieval, label fusion, and offline evaluation"
    )
    doc.add_page_break()

    # 1. Overview
    add_heading(doc, "1. Executive Overview")
    add_para(
        doc,
        "MEDAGENT-X v2 is a clean rebuild of the chest X-ray agentic pipeline on the "
        "architecture-v2 branch. The system ingests MIMIC-CXR / CheXpert-aligned data from "
        "Redivis, builds a label-enriched patient cohort, applies a DICOM quality gate, "
        "trains a fine-tuned RAD-DINO vision model, indexes train studies for image-similarity "
        "retrieval, fuses vision predictions with retrieved report evidence in a deterministic "
        "gray-zone policy, and evaluates everything offline with a structured Judge. The primary "
        "cohort artifact is cohort_balanced_v1 (~1,694 patients, 5,299 studies, 6,130 views)."
    )
    add_para(doc, "Locked pipeline architecture:")
    add_bullets(
        doc,
        [
            "OFFLINE: data → labels → quality → splits → vision training → retrieval index",
            "INFERENCE: Vision → Retrieval → Label Fusion → (optional Report Writer — not yet built)",
            "EVAL: offline Judge only (never in the live inference graph)",
        ],
    )

    # 2. Data acquisition
    add_heading(doc, "2. Phase 1 — Data Acquisition & Label Gating")
    add_para(
        doc,
        "The first phase pulls metadata and imaging from Redivis and restricts the cohort to "
        "studies that have usable CheXpert-style labels. This prevents training or evaluation on "
        "unlabeled views and keeps the data contract auditable."
    )
    add_para(doc, "CLI: medagentx.cli.fetch_stage_a_data")
    add_heading(doc, "Decisions", 2)
    add_bullets(
        doc,
        [
            "Source: Redivis MIMIC-CXR train split metadata + DICOM index + findings_fixed.json",
            "Label-gated cohort: only download DICOMs for studies with parseable findings",
            "Artifacts written under v2/artifacts/cohort/ (or a custom --output-root)",
            "Paginated SQL fetches to respect Redivis response size limits",
        ],
    )
    add_heading(doc, "Implementation", 2)
    add_numbered(
        doc,
        [
            "Fetch eligible DICOM row metadata from Redivis",
            "Download and cache findings_fixed.json",
            "Apply label-gated cohort rules (medagentx.data.cohort)",
            "Download only eligible DICOM files to dicom_train/",
            "Write eligible_dicom_rows.csv and download status summaries",
        ],
    )
    add_para(doc, "Key outputs: eligible_dicom_rows.csv, dicom_train/, findings artifacts")

    # 3. Balanced cohort
    add_heading(doc, "3. Phase 2 — Label-Enriched Balanced Cohort")
    add_para(
        doc,
        "The balanced cohort step selects a label-aware subset of patients and studies so that "
        "rare diseases are represented well enough for training and evaluation. Policy version: "
        "label_enriched_cohort_policy_v2."
    )
    add_para(doc, "CLI: medagentx.cli.build_balanced_cohort")
    add_heading(doc, "Decisions", 2)
    add_bullets(
        doc,
        [
            "Default output: v2/artifacts/cohort_balanced_v1",
            "Max 4 studies per patient (reduces redundant downloads)",
            "Negative-to-positive ratio 2:1 for study selection",
            "Pre-quality positive targets buffered above post-quality targets to absorb failed views",
            "Post-quality targets: train 200 / val 50 / test 50 positive studies (approximate goals)",
            "Can reuse DICOMs from an earlier cohort download (hardlink/copy)",
        ],
    )
    add_heading(doc, "Implementation", 2)
    add_numbered(
        doc,
        [
            "Fetch and label-gate rows (same Redivis flow as Phase 1)",
            "Run enriched cohort selection (medagentx.data.balanced_select)",
            "Download or reuse DICOMs for selected studies",
            "Build preliminary study_label_table for the pool",
            "Write cohort manifest and selection summaries",
        ],
    )

    # 4. Study labels
    add_heading(doc, "4. Phase 3 — Study-Level Label Tables")
    add_para(
        doc,
        "CheXpert labels are parsed from findings, aggregated to study level, and written in a "
        "wide auditable schema. Training uses 12 disease heads; Support Devices and No Finding are "
        "handled under separate label policy rules."
    )
    add_para(doc, "CLI: medagentx.cli.build_study_labels")
    add_heading(doc, "Decisions", 2)
    add_bullets(
        doc,
        [
            "12 disease labels (DISEASE_LABELS): Atelectasis through Pleural Other",
            "Label statuses: present, absent, uncertain, unmentioned (LabelStatus enum)",
            "CheXpert training policy v1: U-mask supervision (uncertain labels masked, not trained as negative)",
            "Multi-view aggregation with conflict flags when views disagree",
            "Representative DICOM path per study: frontal view preferred",
        ],
    )
    add_heading(doc, "Implementation", 2)
    add_numbered(
        doc,
        [
            "Join eligible rows to findings index",
            "build_study_label_bundle per study_key",
            "Flatten to study_label_table.csv with status_*, training_target_*, training_mask_* columns",
        ],
    )

    # 5. Quality gate
    add_heading(doc, "5. Phase 4 — DICOM Quality Gate")
    add_para(
        doc,
        "Before splitting and training, each view is scored on robust cohort-relative image "
        "metrics. Failed views are excluded; borderline views may pass with a technical warning."
    )
    add_para(doc, "CLI: medagentx.cli.run_quality_gate")
    add_heading(doc, "Decisions", 2)
    add_bullets(
        doc,
        [
            "Policy: dicom_quality_policy_v1",
            "Metrics: intensity mean/std, contrast, entropy, sharpness, noise proxies",
            "Robust z-scores vs cohort reference (MAD / IQR fallbacks)",
            "Decisions: pass, review_with_technical_warning, fail",
            "Only quality-passed views proceed to splits and vision training",
        ],
    )
    add_heading(doc, "Implementation", 2)
    add_numbered(
        doc,
        [
            "Compute per-view quality metrics from DICOM pixels",
            "Apply policy thresholds (WARNING_Z=3, FAIL_Z=6)",
            "Write quality/eligible_dicom_rows.csv and per-view decisions",
        ],
    )

    # 6. Finalize cohort
    add_heading(doc, "6. Phase 5 — Cohort Finalization")
    add_para(
        doc,
        "Finalization reconciles quality-filtered views with label tables and produces the "
        "canonical cohort artifacts used by splits and training."
    )
    add_para(doc, "CLI: medagentx.cli.finalize_balanced_cohort")
    add_bullets(
        doc,
        [
            "Merge quality-passed views with study labels",
            "Write final study_label_table.csv at cohort root",
            "Ensure view/split inputs are consistent for downstream CLIs",
        ],
    )

    # 7. Splits
    add_heading(doc, "7. Phase 6 — Patient-Level Splits")
    add_para(
        doc,
        "Train/validation/test assignment is at the patient level to prevent leakage across "
        "studies from the same person."
    )
    add_para(doc, "CLI: medagentx.cli.build_patient_splits")
    add_heading(doc, "Decisions", 2)
    add_bullets(
        doc,
        [
            "Policy: patient_split_policy_v1",
            "Ratios: 70% train / 15% val / 15% test",
            "Random seed: 42 (reproducible)",
            "Unit for vision training and retrieval indexing: study_key",
            "Final cohort split counts: train 3,538 studies / val 855 / test 906",
        ],
    )
    add_heading(doc, "Implementation", 2)
    add_numbered(
        doc,
        [
            "Assign each deid_patient_id to exactly one split",
            "Propagate split to all views in splits/view_splits.csv",
            "No patient appears in more than one split",
        ],
    )

    # 8. Vision training
    add_heading(doc, "8. Phase 7 — Vision Model Training (RAD-DINO)")
    add_para(
        doc,
        "The vision backend fine-tunes microsoft/rad-dino with a multi-label head over 12 "
        "diseases. Study-level prediction pools embeddings from all quality-passed views in a study."
    )
    add_para(doc, "CLI: medagentx.cli.train_raddino")
    add_heading(doc, "Decisions", 2)
    add_bullets(
        doc,
        [
            "Backend ID: raddino_finetuned_v1",
            "Backbone: microsoft/rad-dino; last 2 transformer blocks trainable",
            "Loss: masked binary cross-entropy on U-mask supervised cells only",
            "Checkpoint selection: best validation masked macro AUROC",
            "Per-label present thresholds tuned on validation (F1-oriented search)",
            "Default batch size 2 with gradient accumulation 4 (effective batch 8)",
            "Early stopping patience: 5 epochs",
        ],
    )
    add_heading(doc, "Implementation", 2)
    add_numbered(
        doc,
        [
            "StudyBatchCollator pools views → study_embeddings → classifier logits",
            "train_with_validation saves best_checkpoint.pt under vision/raddino_finetuned_v1/",
            "Thresholds and model metadata stored in checkpoint JSON",
        ],
    )
    add_heading(doc, "Training results (cohort_balanced_v1)", 2)
    add_bullets(
        doc,
        [
            "Test macro AUROC: 0.796",
            "Test macro F1 (masked training metric): 0.477",
            "Checkpoint: vision/raddino_finetuned_v1/best_checkpoint.pt",
        ],
    )

    # 9. Vision inference
    add_heading(doc, "9. Phase 8 — Vision Inference")
    add_para(
        doc,
        "At inference, the fine-tuned backend emits study-level probabilities and binary "
        "present/absent statuses per disease using validation-tuned thresholds."
    )
    add_para(doc, "CLI: medagentx.cli.predict_raddino")
    add_bullets(
        doc,
        [
            "Output contract: probability_{label}, threshold_{label}, status_{label} per study",
            "Extended API: predict_study_outputs() also returns study_embeddings for retrieval",
            "One DICOM forward pass provides both classification and retrieval query vectors",
        ],
    )

    # 10. Retrieval
    add_heading(doc, "10. Phase 9 — Retrieval Index (Offline)")
    add_para(
        doc,
        "Similar historical cases are indexed from train studies only. At inference, the current "
        "study's image embedding queries this index; report text from neighbors supplies fusion evidence."
    )
    add_para(doc, "CLI: medagentx.cli.build_retrieval_index")
    add_heading(doc, "Decisions", 2)
    add_bullets(
        doc,
        [
            "Policy: raddino_image_retrieval_policy_v1",
            "Similarity: fine-tuned RAD-DINO study_embeddings (image-only, not report text)",
            "Index corpus: train split only (3,538 studies) — no val/test leakage",
            "Study-level: one embedding and one index row per study_key",
            "Stored payload: section_findings + section_impression (no GT labels in index)",
            "Vector store: ChromaDB, cosine distance, collection medagentx_train_studies_v1",
            "Query defaults: top_k=5; exclude same study and same patient",
            "Embeddings cache: study_embeddings.npz for faster rebuilds",
            "Progress logging on long embed loops",
        ],
    )
    add_heading(doc, "Implementation", 2)
    add_numbered(
        doc,
        [
            "extract_study_embeddings from train studies",
            "build_index_records + write_chroma_index",
            "write_index_manifest.csv for audit",
            "Inference helper: reasoning/retrieve.py → retrieve_similar_reports()",
        ],
    )

    # 11. Label fusion
    add_heading(doc, "11. Phase 10 — Label Fusion")
    add_para(
        doc,
        "Label Fusion adjusts vision predictions only in a probability gray zone, using "
        "deterministic keyword/negation counts from retrieved reports. No LLM is used in fusion."
    )
    add_para(doc, "Package: medagentx.reasoning")
    add_heading(doc, "Decisions", 2)
    add_bullets(
        doc,
        [
            "Policy: deterministic_gray_zone_fusion_v1",
            "Mechanism: rules-only (no LLM in fusion; Report Writer deferred)",
            "Gray zone: |probability − threshold| ≤ 0.15",
            "Strong zone (outside gray): fusion never overrides vision",
            "Promotion: gray absent → present if ≥3 positive mentions and 0 negative",
            "Demotion: gray present → uncertain if negatives > positives; → absent if 0 positive and ≥3 negative",
            "All 12 diseases share the same policy",
            "Fusion outputs 12 disease statuses only (No Finding not inferred at fusion time)",
            "Reuse vision study_embeddings for retrieval query (single DICOM pass)",
            "Implementation rewritten for v2 (v1 patterns only; not copied monolith)",
        ],
    )
    add_heading(doc, "Implementation", 2)
    add_numbered(
        doc,
        [
            "reasoning/constants.py — locked policy knobs",
            "reasoning/mentions.py — LABEL_TERMS, NEGATION_TERMS, count_retrieval_mentions()",
            "reasoning/fuse.py — fuse_study_labels(), gray-zone gates",
            "reasoning/vision_adapter.py — VisionStudyOutput → fusion inputs",
            "reasoning/retrieve.py — Chroma query at inference",
        ],
    )

    # 12. Fusion eval
    add_heading(doc, "12. Phase 11 — Fusion Evaluation CLI")
    add_para(
        doc,
        "The fusion evaluation script runs the full inference stack on val or test, writes "
        "auditable CSVs, and scores vision-only vs fusion with the offline Judge."
    )
    add_para(doc, "CLI: medagentx.cli.run_fusion_eval")
    add_numbered(
        doc,
        [
            "predict_study_outputs() on chosen split (val/test only)",
            "Per study: retrieve top-5 train neighbors → fuse_study_labels()",
            "Write vision_study_predictions.csv and fusion_label_predictions.csv",
            "Load ground truth from study_label_table.csv",
            "Run Judge: vision_full, fusion_full, vision_gray_zone, fusion_gray_zone",
            "Write judge_summary.json (macro/micro precision, recall, F1, coverage)",
        ],
    )

    # 13. Judge
    add_heading(doc, "13. Phase 12 — Offline Judge")
    add_para(
        doc,
        "The Judge is not a LangGraph agent. It is a deterministic evaluation library that "
        "compares predicted LabelStatus values to frozen ground truth and computes metrics."
    )
    add_heading(doc, "Decisions", 2)
    add_bullets(
        doc,
        [
            "Policy: judge_metric_policy_v1",
            "Ground truth: study_label_table.csv (chexpert_weak_gt_policy_v1)",
            "Compares structured statuses only — does not parse reports at eval time",
            "TP/TN/FP/FN on definite present/absent GT; uncertain/unmentioned GT → unresolved or miss_uncertain",
            "Macro metrics: average across 12 disease labels",
            "Never runs in live inference graph",
        ],
    )

    # 14. Results
    add_heading(doc, "14. Results (Test Split, 906 Studies)")
    add_para(
        doc,
        "Full test evaluation on cohort_balanced_v1 with deterministic gray-zone fusion "
        "(margin 0.15, promotion ≥3 positive / 0 negative)."
    )

    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Run"
    hdr[1].text = "Macro F1"
    hdr[2].text = "Macro Precision"
    hdr[3].text = "Macro Recall"

    rows_data = [
        ("vision_full (overall vision baseline)", "0.639", "0.918", "0.494"),
        ("fusion_full (overall fusion — headline)", "0.678", "0.922", "0.539"),
        ("vision_gray_zone", "0.626", "0.875", "0.493"),
        ("fusion_gray_zone", "0.730", "0.887", "0.628"),
    ]
    for name, f1, prec, rec in rows_data:
        row = table.add_row().cells
        row[0].text = name
        row[1].text = f1
        row[2].text = prec
        row[3].text = rec

    doc.add_paragraph()
    add_para(doc, "Additional context:")
    add_bullets(
        doc,
        [
            "Coverage ~15%: only ~15% of study–label cells have definite present/absent GT (CheXpert uncertain/unmentioned excluded)",
            "Fusion changed 389 / 10,872 cells (3.6%); 282 / 906 studies had ≥1 change",
            "372 promotions (absent→present) vs 17 demotions (present→uncertain)",
            "Top promoted labels: Atelectasis, Pleural Effusion, Edema, Consolidation, Lung Opacity",
            "Gray-zone cells: ~23% of all label cells; fusion improved gray-zone macro F1 by +0.10",
            "Micro F1 fusion_full: 0.708 (pooled across scoreable cells)",
        ],
    )

    add_heading(doc, "Interpretation", 2)
    add_para(
        doc,
        "The agent stack demonstrates measurable value: fusion improves overall test macro F1 "
        "from 0.639 to 0.678 (+0.039) by raising recall (0.494 → 0.539) with minimal precision "
        "loss. The largest gain appears in the gray zone (+0.10 F1), which is exactly where the "
        "retrieval+fusion policy was designed to operate. The system remains precision-heavy "
        "(~0.92 macro precision), indicating room to improve recall further without collapsing "
        "precision."
    )

    # 15. Three steps
    add_heading(doc, "15. Three Recommended Steps to Exceed 0.70 Macro F1")
    add_para(
        doc,
        "Overall headline metric: fusion_full macro F1 (currently 0.678). Reaching 0.70+ likely "
        "requires both better fusion precision on promotions and a stronger vision backbone, "
        "because ~77% of label cells are outside the gray zone and cannot be changed by fusion."
    )

    add_heading(doc, "Step 1 — Val-set error analysis and fusion rule tuning", 2)
    add_bullets(
        doc,
        [
            "Run run_fusion_eval on val (not test) and label each of the ~372-style promotions as TP/FP/FN",
            "Tighten promotion rules for noisy labels (Pleural Effusion, Lung Opacity): require ≥4 positives, stricter keywords, or similarity-weighted counts",
            "Increase demotions when retrieval is strongly negative (only 17 demotions on test today)",
            "Tune on val; lock rules; run test once",
        ],
    )

    add_heading(doc, "Step 2 — Improve vision recall on weak disease heads", 2)
    add_bullets(
        doc,
        [
            "Re-tune per-label thresholds on validation for macro F1 (not AUROC alone)",
            "Focus on Pneumonia, Consolidation, Fracture — labels where vision recall limits full F1",
            "Consider additional training epochs or class-balanced sampling for rare positives",
            "Vision gains lift all cells including the 77% fusion cannot touch",
        ],
    )

    add_heading(doc, "Step 3 — Smarter retrieval evidence (similarity-weighted fusion)", 2)
    add_bullets(
        doc,
        [
            "Weight keyword mentions by retrieval cosine similarity so closer neighbors count more",
            "Require evidence across multiple neighbors (not one noisy report triggering promotion)",
            "Reduces false promotions from generic terms (effusion, opacity) while keeping gray-zone recall gains",
            "Implement as v2 of fusion policy after Step 1 baseline tuning",
        ],
    )

    # 16. Future
    add_heading(doc, "16. Future Work (Not Yet Implemented)")
    add_bullets(
        doc,
        [
            "LangGraph agent wrapper around Vision → Retrieval → Fusion",
            "LLM Report Writer (uses fused labels as binding facts)",
            "Per-label fusion policies and per_label_metrics.csv in Judge output",
            "Vision-phase per-batch progress logging in run_fusion_eval",
            "Optional No Finding consistency pass after fusion",
        ],
    )

    # Appendix - CLI reference
    add_heading(doc, "Appendix A — CLI Command Reference")
    cli_table = doc.add_table(rows=1, cols=2)
    cli_table.style = "Table Grid"
    cli_table.rows[0].cells[0].text = "Phase"
    cli_table.rows[0].cells[1].text = "Command module"
    cli_rows = [
        ("Data fetch", "python -m medagentx.cli.fetch_stage_a_data"),
        ("Balanced cohort", "python -m medagentx.cli.build_balanced_cohort"),
        ("Study labels", "python -m medagentx.cli.build_study_labels"),
        ("Quality gate", "python -m medagentx.cli.run_quality_gate"),
        ("Finalize cohort", "python -m medagentx.cli.finalize_balanced_cohort"),
        ("Patient splits", "python -m medagentx.cli.build_patient_splits"),
        ("Train RAD-DINO", "python -m medagentx.cli.train_raddino"),
        ("Vision predict", "python -m medagentx.cli.predict_raddino"),
        ("Retrieval index", "python -m medagentx.cli.build_retrieval_index"),
        ("Fusion eval", "python -m medagentx.cli.run_fusion_eval"),
    ]
    for phase, cmd in cli_rows:
        r = cli_table.add_row().cells
        r[0].text = phase
        r[1].text = cmd

    add_heading(doc, "Appendix B — Key Artifact Paths (cohort_balanced_v1)")
    add_bullets(
        doc,
        [
            "eligible_dicom_rows.csv, study_label_table.csv, dicom_train/",
            "splits/view_splits.csv",
            "quality/eligible_dicom_rows.csv",
            "vision/raddino_finetuned_v1/best_checkpoint.pt",
            "retrieval/raddino_train_v1/chroma/, study_embeddings.npz",
            "reasoning/fusion_eval_v1/test/vision_study_predictions.csv",
            "reasoning/fusion_eval_v1/test/fusion_label_predictions.csv",
            "reasoning/fusion_eval_v1/test/judge_summary.json",
        ],
    )

    return doc


def main() -> None:
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc = build_document()
    doc.save(OUTPUT_PATH)
    print(f"Wrote {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
